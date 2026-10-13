#!/usr/bin/env python3
"""
Junior-high word data (New Horizon 1, 2, 3) + school-level labels for every book.

Source: Tokyo Shoseki's "NEW HORIZON 単語・連語データベース" (R7 v2) kept in
data/jhs-source/. Sheet ① (Word List) lists every word the first time it appears (学年 = 1/2/3), with
its unit/part code (単元/パート). Sheet ③ (Word List 外の語) lists extra words from Word Room, Unit
Activity Plus, Stage Activity Plus and Further Reading.

Use:
  python tools/build_jhs.py          merge into the existing data/words.js (+ words.json), idempotent
  (tools/build_words.py calls finalize() too, so a full rebuild keeps NH1-3 and the levels)

finalize(out, report) does two things to the data dict:
  1. adds levels to every book:  level 'es'|'jhs' and years [..]  (ES: Let's Try 1,2 = 3,4; NH5,6 = 5,6; Picture Dictionary = 5,6;
     JHS: New Horizon 1,2,3 = 1st, 2nd, 3rd year)
  2. adds the NH1-3 books/decks/words from the spreadsheet (needs openpyxl), and orders books ES (3..6) then JHS (1..3).
"""
import json, os, re, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "data", "jhs-source", "nh-tango-database.xlsx")

# book id -> (level, school years it is meant for, sort order on the shelf)
LEVELS = {
    "lt1": ("es", [3], 0), "lt2": ("es", [4], 1),
    "nh5": ("es", [5], 2), "nh6": ("es", [6], 3),
    "pd": ("es", [5, 6], 4),
    "nh1": ("jhs", [1], 5), "nh2": ("jhs", [2], 6), "nh3": ("jhs", [3], 7),
}
# The owner's rule: "Let's Try 1, 2 is for ES 3, 4" -> Let's Try 1 = 3rd year, Let's Try 2 = 4th year.

# Tile background for decks that have no picture (matches each cover)
TINTS = {1: "#BFE8AE", 2: "#F9C1B0", 3: "#AEDAF6"}

# Parts of the book that are not numbered units. code prefix -> (deck suffix, label, Japanese title)
EXTRAS = [
    ("SaL", "sal", "Sounds & Letters", "おとと もじ"),
    ("GfC", "gfc", "GfC", "ぶんぽう"),
    ("SA", "sa", "Stage Activity", "ステージ アクティビティ"),
    ("LR", "lr", "Let's Read", "よみもの"),
    ("RLE", "rle", "RLE", "よみもの プラス"),
    ("WR", "wr", "Word Room", "ワード ルーム"),
]
EXTRA_PLUS = ("plus", "Plus", "プラス")   # sheet ③: words that are not in the text's Word List


def clean_english(t):
    """What a student types. Placeholders like '...', '～', '(…ing)' and '[are]' are dropped."""
    t = t.replace("’", "'").replace("［", "[").replace("］", "]").replace("（", "(").replace("）", ")")
    t = t.replace("(s)", "s")
    t = re.sub(r"\[[^\]]*\]|\([^)]*\)", " ", t)
    t = re.sub(r"\.{2,}\w*|[～〜~…]", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    t = re.sub(r"\s+([?!.,])", r"\1", t)
    return t.strip()


def typed_letters(t):
    return re.sub(r"[^a-z0-9]", "", t.lower())


def clean_ja(s):
    s = (s or "").strip()
    s = re.sub(r"^【[^】]*】", "", s)         # grammar tag at the front, e.g. 【不定詞（名詞的）】
    s = re.sub(r"^［[^］]*］", "", s)         # e.g. ［提案・勧誘］
    return s.strip()


def section_of(code, grade):
    """Unit code -> (deck suffix, label, title, unit number or None)."""
    code = str(code or "").strip()
    m = re.match(r"U(\d+)(?:-|$)", code)
    if m:
        n = int(m.group(1))
        return f"u{n}", f"Unit {n}", "", n
    for pre, suf, label, title in EXTRAS:
        if code.startswith(pre):
            return suf, label, title, None
    return None


def read_sheets(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    s1 = [r for r in wb.worksheets[0].iter_rows(min_row=4, values_only=True) if r[0] is not None and r[2]]
    s3 = [r for r in wb.worksheets[2].iter_rows(min_row=3, values_only=True) if r[0] is not None and r[1]]
    return s1, s3


def add_jhs(out, report, path=XLSX):
    try:
        s1, s3 = read_sheets(path)
    except ImportError:
        report.append("! openpyxl not installed; New Horizon 1-3 skipped (pip install openpyxl)")
        return
    words, decks = out["words"], out["decks"]
    # drop any earlier NH1-3 so the merge can be re-run
    out["books"] = [b for b in out["books"] if b["id"] not in ("nh1", "nh2", "nh3")]
    for did in [d for d, v in decks.items() if v["book"] in ("nh1", "nh2", "nh3")]:
        del decks[did]

    per = {1: {}, 2: {}, 3: {}}   # grade -> deck suffix -> deck dict

    def deck_for(grade, suf, label, title, unit):
        bid = f"nh{grade}"
        did = f"{bid}-{suf}"
        if did not in per[grade]:
            per[grade][did] = {"id": did, "book": bid, "unit": unit, "label": label, "title": title,
                               "image": "", "words": [], "sentences": [], "_s": set()}
        return per[grade][did]

    def put(d, text, ja, original):
        letters = typed_letters(text)
        if len(letters) <= 1:
            return
        wid = text.lower()
        if wid not in words:
            words[wid] = {"t": text, "ja": ja, "kana": ""}
        sentence = bool(re.search(r"[.?!]$", text)) and " " in text and not re.search(r"\.{2,}|[～〜…]", original)
        if sentence:
            if text not in d["_s"]:
                d["_s"].add(text); d["sentences"].append({"t": text, "ja": ja})
        if wid not in d["words"]:
            d["words"].append(wid)

    skipped = []
    for r in s1:
        grade, code = r[10], r[12]
        if grade not in per:
            continue
        sec = section_of(code, grade)
        if not sec:
            skipped.append((grade, code, r[2])); continue
        text = clean_english(str(r[2]))
        d = deck_for(grade, *sec)
        put(d, text, clean_ja(r[9]), str(r[2]))
    for r in s3:
        grade = r[4]
        if grade not in per:
            continue
        text = clean_english(str(r[1]))
        d = deck_for(grade, EXTRA_PLUS[0], EXTRA_PLUS[1], EXTRA_PLUS[2], None)
        put(d, text, clean_ja(r[3]), str(r[1]))
    if skipped:
        report.append(f"! JHS rows with an unknown unit code (skipped): {skipped[:8]}")

    def sort_key(d):
        suf = d["id"].split("-", 1)[1]
        if d["unit"] is not None:
            return (0, d["unit"])
        order = [e[1] for e in EXTRAS] + [EXTRA_PLUS[0]]
        return (1, order.index(suf))

    for g in (1, 2, 3):
        bid = f"nh{g}"
        ds = sorted(per[g].values(), key=sort_key)
        # Unit 0 must not use unit=0 (the elementary "All" decks use unit 0 to mean "everything")
        for d in ds:
            if d["unit"] == 0:
                d["unit"] = None
            d.pop("_s", None)
            decks[d["id"]] = d
        out["books"].append({"id": bid, "name": f"New Horizon {g}", "short": f"NH{g}", "cover": f"assets/books/{bid}.jpg",
                             "tint": TINTS[g], "decks": [d["id"] for d in ds]})


def finalize(out, report=None, path=XLSX):
    report = report if report is not None else []
    if os.path.isfile(path):
        add_jhs(out, report, path)
    else:
        report.append(f"- no JHS spreadsheet at {os.path.relpath(path, ROOT)}; New Horizon 1-3 not added")
    for b in out["books"]:
        lv = LEVELS.get(b["id"])
        if lv:
            b["level"], b["years"] = lv[0], lv[1]
    out["books"].sort(key=lambda b: LEVELS.get(b["id"], ("", [], 99))[2])
    out["version"] = max(out.get("version", 1), 2)
    return out


def main():
    jp = os.path.join(ROOT, "data", "words.js")
    txt = open(jp, encoding="utf-8").read()
    out = json.loads(txt[txt.index("{"): txt.rindex("}") + 1])
    report = []
    finalize(out, report)
    with open(os.path.join(ROOT, "data", "words.json"), "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    with open(jp, "w", encoding="utf-8") as f:
        f.write("// Generated by tools/build_words.py + tools/build_jhs.py - do not edit by hand.\n")
        f.write("window.TM_WORDS = ")
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";\n")
    print(f"books {len(out['books'])}  decks {len(out['decks'])}  words {len(out['words'])}")
    for b in out["books"]:
        if b["id"].startswith("nh") and b["id"] in ("nh1", "nh2", "nh3"):
            print(f"  {b['short']} ({b['level']} {b['years']}): " + ", ".join(
                f"{d.split('-',1)[1]}({len(out['decks'][d]['words'])}w/{len(out['decks'][d]['sentences'])}s)" for d in b["decks"]))
    for r in report:
        print(r)


if __name__ == "__main__":
    main()
